import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()

# Standard layouts
blank_slide_layout = prs.slide_layouts[6] 
title_slide_layout = prs.slide_layouts[0]
bullet_slide_layout = prs.slide_layouts[1]

slides_data = [
    {
        "title": "ProofBridge: Cryptographic Skill Verification Network",
        "content": [
            "Stop Trusting What Students Claim. Mathematically Verify What They Can Build.",
            "IIC 3.0 MUJ Hackathon · Theme: EdTech · PS-08 Academia-Industry Collaboration",
            "Team Name: Outliers",
            "Team Members: Akshar Miyani, Lubhanshi Mathur, Lakshita Sharma, Mohammad Faizan"
        ]
    },
    {
        "title": "The Core Problem: Resumes Are Broken",
        "content": [
            "Resume Inflation Crisis: Over 78% of engineering resumes contain exaggerated or unverified claims.",
            "Generative AI CV Flooding: Recruiters spend over 42 hours per technical role filtering AI-generated resume spam.",
            "Degree Detachment: Broad semester letter grades fail to signal modular production competence (SQL, Docker, APIs).",
            "Pedigree Bias: Talented students from regional/tier-2 colleges are filtered out by generic keyword ATS tools."
        ]
    },
    {
        "title": "The Market Gap & Stakeholders",
        "content": [
            "The Missing Link: Existing portals (LinkedIn, Internshala) treat skills as self-typed text strings with zero proof.",
            "The Cryptographic Chain: University Anchor -> GitHub Git Evidence -> Faculty Rubrics -> Role Genomes -> W3C Credentials.",
            "Students: Verifiable Skill Twin & shortest-path employability roadmap.",
            "Faculty Evaluators: Structured rubric queues with < 24h SLA and NAAC/NBA accreditation data.",
            "Industry Employers: Blind vector talent search and 80% reduction in technical screening time.",
            "University Leadership: Real-time curriculum demand gap radar across student cohorts."
        ]
    },
    {
        "title": "The Proposed Solution: ProofBridge Platform",
        "content": [
            "What is ProofBridge? An evidence-based talent intelligence network connecting academia and industry via cryptographic code proofs.",
            "Pillar 1 - Skill Twin: Dynamic digital twin of student competence linked to frozen Git commit digests.",
            "Pillar 2 - Role Genome: Requisitions converted into multi-factor capability vectors with weighted rubric levels.",
            "Pillar 3 - Bridge Me Engine: Shortest-path pathfinder calculating highest-ROI challenges to close eligibility gaps.",
            "Pillar 4 - Faculty Trust Anchor: Accredited human faculty scoring with 0% AI hallucination, outputting W3C JSON-LD credentials."
        ]
    },
    {
        "title": "GitHub Ground Truth & Evidence Verification",
        "content": [
            "Real Code as Ground Truth: Direct analysis of GitHub repositories, AST structures, and test suites (e.g. 2,840 lines analyzed).",
            "Mandatory Student Contribution Defense: Candidates must explain independent problem-solving vs tooling assistance.",
            "AI & Tooling Transparency Disclosure: Tagging independent authoring, AI syntax verification, or autocomplete copilot.",
            "Frozen Revision Locks: Atomic PostgreSQL submission locks with server-calculated SHA-256 digests."
        ]
    },
    {
        "title": "Technology Stack & Deterministic Engine",
        "content": [
            "Frontend: Next.js 14 (App Router), React 18, TypeScript, Tailwind CSS, Framer Motion, Next-Themes.",
            "Backend & Storage: PostgreSQL with connection pooling, immutable revision logs, atomic state synchronization.",
            "W3C Verifiable Credentials: JSON-LD standard proofs, SHA-256 code digests, Did:pb decentralized identifiers.",
            "coverage-v1 Algorithm: Deterministic mathematical scoring formula [Score = Sum(Weight * min(1, Level / Target))]."
        ]
    },
    {
        "title": "Curriculum Intelligence & Blind Talent Radar",
        "content": [
            "University Curriculum Gap Radar: Real-time cohort demand analytics (e.g. 620 students: 142 ready, 213 near-ready, 49% Docker gap).",
            "Proactive Curriculum Adaptation: Enables deans to run targeted bootcamps before placement season begins.",
            "Blind Vector Hiring: Candidates shortlisted purely by mathematical genome match with zero pedigree or demographic bias.",
            "1-Click Verification: Third-party employers verify code digests offline without contacting university administration."
        ]
    },
    {
        "title": "Feasibility, Accreditation & Measurable Impact",
        "content": [
            "Zero Recurring Cost: Open W3C standards and self-hosted PostgreSQL eliminate expensive third-party AI/API bills.",
            "Accreditation Compliance: Automated data export for NAAC Criteria 2 & 5, NBA OBE Criterion 2.6, and NEP 2020 frameworks.",
            "80% Reduction in Technical Screening: Employers interview pre-verified candidates with zero resume reading.",
            "3.2x Higher Interview Conversion: Students with evidence-backed Skill Twins demonstrate proven technical capability."
        ]
    },
    {
        "title": "Team Outliers & Future Vision",
        "content": [
            "Akshar Miyani: Lead Architect, Full-Stack & Cryptography",
            "Lubhanshi Mathur: System Engineering & UX/UI Design",
            "Lakshita Sharma: Academic Rubrics & Accreditation Frameworks",
            "Mohammad Faizan: Backend Pipelines & Infrastructure",
            "",
            "Vision: Empowering every student to be evaluated by what they can build, not by what they can fabricate."
        ]
    }
]

# Path to assets
ppt_dir = os.path.dirname(os.path.abspath(__file__))
logo_muj = os.path.join(ppt_dir, "assets", "muj-logo.png")
logo_iic3 = os.path.join(ppt_dir, "assets", "iic3-logo.png")

for i, slide_data in enumerate(slides_data):
    if i == 0:
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = slide_data["title"]
        subtitle.text = "\n".join(slide_data["content"])
    else:
        slide = prs.slides.add_slide(bullet_slide_layout)
        title = slide.shapes.title
        body = slide.placeholders[1]
        title.text = slide_data["title"]
        
        tf = body.text_frame
        tf.text = slide_data["content"][0]
        for item in slide_data["content"][1:]:
            p = tf.add_paragraph()
            p.text = item
            p.level = 0
            
    # Add logos to slides if available
    try:
        if os.path.exists(logo_muj):
            slide.shapes.add_picture(logo_muj, Inches(0.4), Inches(0.3), width=Inches(1.8))
        if os.path.exists(logo_iic3):
            slide.shapes.add_picture(logo_iic3, Inches(8.2), Inches(0.3), width=Inches(1.4))
    except Exception as e:
        print(f"Notice on slide {i}: {e}")

output_path = os.path.join(ppt_dir, "ProofBridge_Hackathon_Pitch.pptx")
prs.save(output_path)
print(f"Presentation saved successfully at {output_path}")
