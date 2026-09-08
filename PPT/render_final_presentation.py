import os
import json
import subprocess
import pptx
from pptx.util import Inches
from pypdf import PdfWriter, PdfReader

def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    final_ppt_dir = os.path.join(base_dir, "Final_PPT")
    rendered_dir = os.path.join(final_ppt_dir, "rendered")
    os.makedirs(rendered_dir, exist_ok=True)

    manifest_path = os.path.join(final_ppt_dir, "manifest.json")
    if os.path.exists(manifest_path):
        with open(manifest_path, "r", encoding="utf-8") as f:
            manifest = json.load(f)
        slides = manifest.get("playlist", [])
    else:
        slides = [
            "cover.html",
            "broken-triangle.html",
            "trust-chain.html",
            "core-breakthroughs.html",
            "ast-evidence.html",
            "architecture.html",
            "dean-radar.html",
            "feasibility.html",
            "team-roadmap.html"
        ]

    chrome_path = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
    if not os.path.exists(chrome_path):
        chrome_path = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

    png_files = []
    pdf_files = []

    print(f"Step 1: Rendering {len(slides)} slides at 1920x1080 from {final_ppt_dir}...")
    for idx, slide_file in enumerate(slides):
        slide_num = f"{idx+1:02d}"
        slide_path = os.path.join(final_ppt_dir, slide_file)
        file_url = f"file:///{slide_path.replace(os.sep, '/')}"

        out_png = os.path.join(rendered_dir, f"slide_{slide_num}.png")
        out_pdf = os.path.join(rendered_dir, f"slide_{slide_num}.pdf")

        # High-res Screenshot PNG (1920x1080)
        cmd_png = [
            chrome_path,
            "--headless=new",
            "--disable-gpu",
            "--hide-scrollbars",
            "--window-size=1920,1080",
            f"--screenshot={out_png}",
            file_url
        ]
        subprocess.run(cmd_png, check=True, capture_output=True)
        png_files.append(out_png)

        # High-res Vector PDF slide
        cmd_pdf = [
            chrome_path,
            "--headless=new",
            "--disable-gpu",
            "--hide-scrollbars",
            "--print-to-pdf-no-header",
            f"--print-to-pdf={out_pdf}",
            file_url
        ]
        subprocess.run(cmd_pdf, check=True, capture_output=True)
        pdf_files.append(out_pdf)

        print(f"  [OK] Slide {slide_num}: {slide_file} rendered")

    # Step 2: Assemble Single PDF
    print("\nStep 2: Merging slides into master PDF...")
    merger = PdfWriter()
    for pdf in pdf_files:
        reader = PdfReader(pdf)
        for page in reader.pages:
            merger.add_page(page)

    final_pdf_path = os.path.join(final_ppt_dir, "ProofBridge_Pitch_Deck.pdf")
    with open(final_pdf_path, "wb") as f_out:
        merger.write(f_out)
    print(f"  [OK] Final PDF created: {final_pdf_path} ({os.path.getsize(final_pdf_path)} bytes)")

    # Also save to main PPT folder
    with open(os.path.join(base_dir, "ProofBridge_Pitch.pdf"), "wb") as f_out:
        merger.write(f_out)
    with open(os.path.join(base_dir, "ProofBridge_IIC3_Final_Pitch.pdf"), "wb") as f_out:
        merger.write(f_out)

    # Step 3: Build Full-Bleed 16:9 PowerPoint Presentation (.pptx)
    print("\nStep 3: Creating master PowerPoint presentation (.pptx)...")
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for png in png_files:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(png, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))

    final_pptx_path = os.path.join(final_ppt_dir, "ProofBridge_Pitch_Deck.pptx")
    prs.save(final_pptx_path)
    print(f"  [OK] Final PPTX created: {final_pptx_path} ({os.path.getsize(final_pptx_path)} bytes)")

    # Also save to main PPT folder so all links and downloads stay synchronized
    prs.save(os.path.join(base_dir, "ProofBridge_Pitch.pptx"))
    prs.save(os.path.join(base_dir, "ProofBridge_Hackathon_Pitch.pptx"))
    prs.save(os.path.join(base_dir, "ProofBridge_IIC3_Final_Pitch.pptx"))

    print("\nALL PRESENTATION DELIVERABLES GENERATED SUCCESSFULLY!")

if __name__ == "__main__":
    main()
